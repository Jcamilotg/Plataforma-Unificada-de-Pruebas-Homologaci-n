import os
import sqlite3
import subprocess
import threading
import time
import platform
import csv
import io
from datetime import datetime
from flask import Flask, render_template_string, jsonify, request, Response

# Librerías para generación de Excel e incrustación de gráficos
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt

app = Flask(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ACS_DB_PATH = os.path.join(BASE_DIR, "acs_database.db")
SNMP_DB_PATH = os.path.join(BASE_DIR, "snmp_cache.db")

# Caché global en memoria para evitar re-ejecutar el Walk al descargar
GLOBAL_WALK_CACHE = {"ip": "", "oid": "", "content": ""}

# Control de equipos en monitoreo activo
MONITORED_TARGETS = set()
MONITORED_LOCK = threading.Lock()


def get_acs_db():
    conn = sqlite3.connect(ACS_DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode=WAL;")
    except Exception:
        pass
    return conn


def get_snmp_db():
    conn = sqlite3.connect(SNMP_DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode=WAL;")
    except Exception:
        pass
    return conn


def init_sentinel_db():
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS snmp_cache (
            host_name TEXT PRIMARY KEY,
            ip_address TEXT,
            sys_descr TEXT DEFAULT 'N/A',
            uptime TEXT DEFAULT 'N/A',
            serial_number TEXT DEFAULT 'N/A',
            temperature TEXT DEFAULT 'N/A',
            status TEXT DEFAULT 'Unknown',
            last_check DATETIME DEFAULT CURRENT_TIMESTAMP
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS snmp_services_cache (
            host_name TEXT,
            service_name TEXT,
            oid TEXT,
            value TEXT,
            last_check DATETIME DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(host_name, service_name)
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS device_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            host_name TEXT,
            status TEXT,
            timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS device_drops (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            host_name TEXT,
            drop_time DATETIME DEFAULT CURRENT_TIMESTAMP,
            recovery_time DATETIME,
            duration_seconds INTEGER
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS monitored_targets (
            target_name TEXT PRIMARY KEY
        );
    """)
    conn.commit()

    # Cargar los equipos en monitoreo persistidos previamente
    cursor.execute("SELECT target_name FROM monitored_targets")
    rows = cursor.fetchall()
    with MONITORED_LOCK:
        MONITORED_TARGETS.clear()
        for row in rows:
            MONITORED_TARGETS.add(row["target_name"])

    conn.close()


init_sentinel_db()

SNMP_HOSTS = [
    {"name": "Internet", "ip": "8.8.8.8", "community": "private", "type": "ping"},
    {"name": "Servidor_PPAL (MikroTik)", "ip": "192.168.35.1", "community": "private", "type": "ping"},
    {"name": "CMTS", "ip": "192.168.35.4", "community": "private", "type": "ping"},
    {"name": "OLT", "ip": "192.168.35.10", "community": "private", "type": "ping"},
    {"name": "Switch", "ip": "192.168.35.6", "community": "private", "type": "ping"},
    {"name": "ACS", "ip": "192.168.35.9", "community": "private", "type": "ping"},
    {"name": "SIP_Server", "ip": "192.168.35.9", "community": "private", "type": "ping"},
    {"name": "Cable_Modem_2", "ip": "172.16.6.2", "community": "private", "type": "cm"},
    {"name": "Cable_Modem_3", "ip": "172.16.6.3", "community": "private", "type": "cm"},
    {"name": "Cable_Modem_4", "ip": "172.16.6.4", "community": "private", "type": "cm"},
    {"name": "Cable_Modem_5", "ip": "172.16.6.5", "community": "private", "type": "cm"},
    {"name": "Cable_Modem_6", "ip": "172.16.6.6", "community": "private", "type": "cm"},
    {"name": "Cable_Modem_7", "ip": "172.16.6.7", "community": "private", "type": "cm"},
]

CABLE_MODEM_SERVICES = [
    ("Monitor_System_Uptime", "1.3.6.1.2.1.1.3.0"),
    ("Monitor_SysDescr", "1.3.6.1.2.1.1.1.0"),
    ("WiFi_DFS_Channel_2disable", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.2.1.24.10100"),
    ("EMTA_SIP_0xE0_NCS_0xC0", "1.3.6.1.4.1.4413.2.2.2.1.2.98.1.1.1.11.0"),
    ("EMTA_2_File_1_TR104", "1.3.6.1.4.1.1038.28.1.1.142.0"),
    ("TR69_URL_ACS", "1.3.6.1.4.1.1038.28.1.8.4.2.0"),
    ("Monitor_SwBase_Chipset", "1.3.6.1.4.1.4413.2.2.2.1.9.1.1.8.0"),
    ("Login_Password_Of_Day", "1.3.6.1.4.1.4413.2.2.2.1.1.3.100.0"),
    ("Monitor_Config_File", "1.3.6.1.2.1.69.1.4.5.0"),
    ("Login_HttpAdminID", "1.3.6.1.4.1.4413.2.2.2.1.1.3.1.0"),
    ("Login_HttpAdminPassword", "1.3.6.1.4.1.4413.2.2.2.1.1.3.2.0"),
    ("Login_HttpUserID", "1.3.6.1.4.1.4413.2.2.2.1.1.3.3.0"),
    ("Login_HttpUserPassword", "1.3.6.1.4.1.4413.2.2.2.1.1.3.4.0"),
    ("Login_RemoteAccess_Enable", "1.3.6.1.4.1.1038.28.1.1.17.0"),
    ("Login_sshServerControl", "1.3.6.1.4.1.4413.2.2.2.1.1.4.4.0"),
    ("Login_sshUserName", "1.3.6.1.4.1.4413.2.2.2.1.1.4.2.0"),
    ("Login_sshPassword", "1.3.6.1.4.1.4413.2.2.2.1.1.4.3.0"),
    ("WiFi_SSID_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.2.1.1.3.10001"),
    ("WiFi_SSID_5G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.2.1.1.3.10101"),
    ("WiFi_Password_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.2.3.4.1.2.10001"),
    ("WiFi_Password_5G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.2.3.4.1.2.10101"),
    ("WiFi_CH_selected_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.2.1.2.10000"),
    ("WiFi_CH_selected_5G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.2.1.2.10100"),
    ("WiFi_CH_Operational_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.2.1.43.10000"),
    ("WiFi_CH_Operational_5G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.2.1.43.10100"),
    ("WiFi_Security_Type_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.2.1.1.4.10001"),
    ("WiFi_Bandwidth_2G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.6.1.1.10000"),
    ("WiFi_Bandwidth_5G", "1.3.6.1.4.1.4413.2.2.2.1.18.1.1.6.1.1.10100"),
    ("Monitor_RG_Mode", "1.3.6.1.4.1.4413.2.2.2.1.7.1.4.0"),
    ("WiFi_Plume_Enabled", "1.3.6.1.4.1.1038.28.1.7.29.1.0"),
    ("Monitor_Temperature", "1.3.6.1.4.1.1038.28.1.1.31.0"),
    ("Login_RemoteAccess_port", "1.3.6.1.4.1.1038.28.1.1.20.0"),
    ("Login_RemoteAccess_User", "1.3.6.1.4.1.1038.28.1.1.18.0"),
    ("Login_RemoteAccess_Password", "1.3.6.1.4.1.1038.28.1.1.19.0"),
    ("Monitor_SerialNumber", "1.3.6.1.2.1.69.1.1.4.0"),
    ("MAC_CM", "1.3.6.1.2.1.2.2.1.6.2"),
    ("CM_IPAddress", "1.3.6.1.4.1.1038.28.1.1.7.0")
]


def check_ping(ip):
    try:
        param = "-n" if platform.system().lower() == "windows" else "-c"
        command = ["ping", param, "1", "-W", "1", ip]
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        return result.returncode == 0
    except Exception:
        return False


def query_snmp(ip, community, oid):
    try:
        cmd = ["snmpget", "-v2c", "-c", community, "-t", "2", "-r", "1", ip, oid]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=2)
        if result.returncode == 0:
            output = result.stdout.strip()
            if "=" in output:
                val = output.split("=", 1)[1].strip()
                if ":" in val:
                    val = val.split(":", 1)[1].strip()
                return val.replace('"', '')
    except Exception:
        pass
    return "N/A"


def process_device_status(cursor, host_key, status):
    cursor.execute("""
        SELECT id, drop_time FROM device_drops 
        WHERE host_name = ? AND recovery_time IS NULL 
        ORDER BY id DESC LIMIT 1
    """, (host_key,))
    open_drop = cursor.fetchone()

    if status == "Offline":
        if not open_drop:
            cursor.execute("""
                INSERT INTO device_drops (host_name, drop_time) 
                VALUES (?, datetime('now', 'localtime'))
            """, (host_key,))
    elif status == "Online":
        if open_drop:
            cursor.execute("""
                UPDATE device_drops
                SET recovery_time = datetime('now', 'localtime'),
                    duration_seconds = CAST((strftime('%s', 'now', 'localtime') - strftime('%s', drop_time)) AS INTEGER)
                WHERE id = ?
            """, (open_drop["id"],))

    cursor.execute("""
        INSERT INTO device_history (host_name, status, timestamp) 
        VALUES (?, ?, datetime('now', 'localtime'))
    """, (host_key, status))


def snmp_poller_background():
    while True:
        conn = get_snmp_db()
        cursor = conn.cursor()

        for host in SNMP_HOSTS:
            ip = host["ip"]
            comm = host["community"]
            name = host["name"]
            htype = host["type"]

            with MONITORED_LOCK:
                is_monitored = name in MONITORED_TARGETS

            if not is_monitored:
                continue

            if htype == "ping":
                is_up = check_ping(ip)
                status = "Online" if is_up else "Offline"
                sys_descr = "Infraestructura (Monitoreo por Ping)"
                uptime, serial, temp = "N/A", "N/A", "N/A"
            else:
                is_reachable = check_ping(ip)
                if is_reachable:
                    sys_descr = query_snmp(ip, comm, "1.3.6.1.2.1.1.1.0")
                    uptime = query_snmp(ip, comm, "1.3.6.1.2.1.1.3.0")
                    serial = query_snmp(ip, comm, "1.3.6.1.2.1.69.1.1.4.0")
                    temp = query_snmp(ip, comm, "1.3.6.1.4.1.1038.28.1.1.31.0")
                    status = "Online" if (sys_descr != "N/A" or uptime != "N/A") else "Offline"
                else:
                    sys_descr, uptime, serial, temp = "N/A", "N/A", "N/A", "N/A"
                    status = "Offline"

            process_device_status(cursor, name, status)

            cursor.execute("""
                INSERT INTO snmp_cache (host_name, ip_address, sys_descr, uptime, serial_number, temperature, status, last_check)
                VALUES (?, ?, ?, ?, ?, ?, ?, datetime('now', 'localtime'))
                ON CONFLICT(host_name) DO UPDATE SET
                    ip_address=EXCLUDED.ip_address,
                    sys_descr=EXCLUDED.sys_descr,
                    uptime=EXCLUDED.uptime,
                    serial_number=EXCLUDED.serial_number,
                    temperature=EXCLUDED.temperature,
                    status=EXCLUDED.status,
                    last_check=datetime('now', 'localtime');
            """, (name, ip, sys_descr, uptime, serial, temp, status))

            if htype == "cm" and status == "Online":
                for svc_name, oid_val in CABLE_MODEM_SERVICES:
                    val = query_snmp(ip, comm, oid_val)
                    cursor.execute("""
                        INSERT INTO snmp_services_cache (host_name, service_name, oid, value, last_check)
                        VALUES (?, ?, ?, ?, datetime('now', 'localtime'))
                        ON CONFLICT(host_name, service_name) DO UPDATE SET
                            oid=EXCLUDED.oid,
                            value=EXCLUDED.value,
                            last_check=datetime('now', 'localtime');
                    """, (name, svc_name, oid_val, val))

        try:
            acs_conn = get_acs_db()
            acs_cursor = acs_conn.cursor()
            acs_cursor.execute("SELECT serial_number, status FROM devices")
            onts = acs_cursor.fetchall()
            acs_conn.close()

            for ont in onts:
                serial = ont["serial_number"]
                with MONITORED_LOCK:
                    is_monitored = serial in MONITORED_TARGETS

                if not is_monitored:
                    continue

                status = ont["status"] or "Offline"
                process_device_status(cursor, serial, status)
        except Exception:
            pass

        conn.commit()
        conn.close()
        time.sleep(10)


threading.Thread(target=snmp_poller_background, daemon=True).start()

HTML_DASHBOARD = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>CPE Sentinel - Monitoreo Proactivo Unificado</title>
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css">
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        body { background-color: #0f172a; color: #f8fafc; font-family: sans-serif; }
        .card { background-color: #1e293b; border: 1px solid #334155; color: #f8fafc; }
        .table { color: #f8fafc; }
        .table-hover tbody tr:hover { background-color: #334155; color: #fff; }
        .badge-danger-custom { background-color: #ef4444; color: white; }
        .badge-success-custom { background-color: #10b981; color: white; }
        .card-detail-section { border: 1px solid #334155; border-radius: 8px; padding: 14px; background-color: #0f172a; height: 100%; }
        .card-detail-section h6 { font-size: 0.95rem; color: #38bdf8; font-weight: 700; margin-bottom: 10px; border-bottom: 1px solid #334155; padding-bottom: 6px; }
        .detail-item { display: flex; justify-content: space-between; font-size: 0.85rem; margin-bottom: 6px; }
        .detail-label { color: #94a3b8; font-weight: 500; }
        .detail-val { font-weight: 600; color: #f8fafc; text-align: right; }
        .modal-content { background-color: #1e293b; color: #f8fafc; border: 1px solid #334155; }
        .section-header { color: #38bdf8; font-weight: 700; border-left: 4px solid #38bdf8; padding-left: 10px; margin-bottom: 15px; margin-top: 35px; }
    </style>
</head>
<body>
<div class="container-fluid py-4 px-4">
    <div class="d-flex justify-content-between align-items-center mb-4">
        <div>
            <h2>🛡️ CPE Sentinel — Panel Unificado Proactivo</h2>
            <p class="text-secondary mb-0">Visualización simultánea de ONTs (TR-069), Infraestructura y Cable Módems</p>
        </div>
        <div>
            <span class="badge bg-primary fs-6 me-2">Puerto 9035</span>
            <button class="btn btn-sm btn-outline-warning me-2" onclick="reiniciarMonitoreoGlobal()">🧹 Reiniciar Histórico Global</button>
            <button class="btn btn-sm btn-outline-light" onclick="cargarTodo()">🔄 Actualizar Todo</button>
        </div>
    </div>

    <div class="row g-4 mb-3">
        <div class="col-md-4">
            <div class="card p-3 shadow-sm">
                <h6 class="text-muted">Total ONTs TR-069</h6>
                <h3 id="total_devs" class="fw-bold">0</h3>
            </div>
        </div>
        <div class="col-md-4">
            <div class="card p-3 shadow-sm">
                <h6 class="text-muted">Alertas Fibra Óptica (&lt; -27 dBm)</h6>
                <h3 id="optic_alerts" class="fw-bold text-danger">0</h3>
            </div>
        </div>
        <div class="col-md-4">
            <div class="card p-3 shadow-sm">
                <h6 class="text-muted">ONTs Offline</h6>
                <h3 id="offline_devs" class="fw-bold text-warning">0</h3>
            </div>
        </div>
    </div>

    <div class="section-header">📡 Cuadro 1: Dispositivos TR-069 (ONTs)</div>
    <div class="card shadow-sm">
        <div class="card-header bg-dark fw-bold">Inventario de Hardware y Estado de Suscriptores</div>
        <div class="card-body p-0 table-responsive">
            <table class="table table-hover align-middle mb-0">
                <thead class="table-dark">
                    <tr>
                        <th>Estado</th>
                        <th>Serial</th>
                        <th>Modelo</th>
                        <th>Hardware</th>
                        <th>Firmware</th>
                        <th>IP Datos</th>
                        <th>Potencia Rx (Fibra)</th>
                        <th>Último Inform</th>
                        <th class="text-end">Acción</th>
                    </tr>
                </thead>
                <tbody id="tr069TableBody">
                    <tr><td colspan="9" class="text-center py-4">Cargando ONTs...</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <div class="section-header">⚙️ Cuadro 2: Equipos de Infraestructura y Cable Módems</div>
    <div class="card shadow-sm mb-4">
        <div class="card-header bg-dark fw-bold">Monitoreo de Redes, Servidores, OLT y Cable Módems</div>
        <div class="card-body p-0 table-responsive">
            <table class="table table-hover align-middle mb-0">
                <thead class="table-dark">
                    <tr>
                        <th>Estado</th>
                        <th>Nombre del Host</th>
                        <th>Dirección IP</th>
                        <th>Descripción del Sistema / Método</th>
                        <th>Número de Serie</th>
                        <th>Temperatura</th>
                        <th>Uptime</th>
                        <th>Último Check</th>
                        <th class="text-end">Acciones</th>
                    </tr>
                </thead>
                <tbody id="snmpTableBody">
                    <tr><td colspan="9" class="text-center py-4">Cargando equipos...</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <!-- SECCIÓN INTERACTIVA DE COMANDOS SNMP CUSTOM CON WALK COMPLETO Y PRESETS -->
    <div class="section-header">🖥️ Consola Interactiva SNMP (GET / SET / WALK)</div>
    <div class="card shadow-sm mb-4 p-3">
        <div class="mb-3 p-2 bg-dark rounded border border-secondary">
            <span class="small text-info fw-bold me-2">📌 Accesos Rápidos OIDs Principales:</span>
            <button type="button" class="btn btn-sm btn-outline-info me-1 mb-1" onclick="setOidPreset('1.3.6.1.2')">mgmt (1.3.6.1.2)</button>
            <button type="button" class="btn btn-sm btn-outline-info me-1 mb-1" onclick="setOidPreset('1.3.6.1.4.1.4491')">cableLabs (1.3.6.1.4.1.4491)</button>
            <button type="button" class="btn btn-sm btn-outline-info me-1 mb-1" onclick="setOidPreset('1.3.6.1.4.1.4413')">broadcom (1.3.6.1.4.1.4413)</button>
            <button type="button" class="btn btn-sm btn-outline-info me-1 mb-1" onclick="setOidPreset('1.3.6.1.4.1.1038')">sagemDr (1.3.6.1.4.1.1038)</button>
        </div>
        <form id="snmpCustomForm" onsubmit="ejecutarSnmpCustom(event)">
            <div class="row g-3">
                <div class="col-md-3">
                    <label class="form-label small text-muted fw-bold">Seleccionar Dispositivo:</label>
                    <select id="snmpTargetSelect" class="form-select bg-dark text-white border-secondary" required>
                        <option value="">Cargando dispositivos...</option>
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label small text-muted fw-bold">Operación:</label>
                    <select id="snmpOpType" class="form-select bg-dark text-white border-secondary" onchange="onOpTypeChange()">
                        <option value="GET">SNMP GET</option>
                        <option value="SET">SNMP SET</option>
                        <option value="WALK">SNMP BULKWALK (Árbol Completo Rápido)</option>
                    </select>
                </div>
                <div class="col-md-3">
                    <label class="form-label small text-muted fw-bold">OID SNMP:</label>
                    <input type="text" id="snmpOidInput" class="form-control bg-dark text-white border-secondary font-monospace" placeholder="Ej: 1.3.6.1.2.1.1.3.0" required>
                </div>
                <div class="col-md-2" id="snmpValContainer" style="display: none;">
                    <label class="form-label small text-muted fw-bold">Valor / Tipo (Ej: i 1, s texto):</label>
                    <input type="text" id="snmpValueInput" class="form-control bg-dark text-white border-secondary font-monospace" placeholder="i 1">
                </div>
                <div class="col-md-2 d-flex align-items-end">
                    <button type="submit" class="btn btn-primary w-100 fw-bold">🚀 Ejecutar Command</button>
                </div>
            </div>
        </form>
        <div id="snmpResultContainer" class="mt-3 p-3 bg-dark border border-secondary rounded" style="display: none;">
            <div class="d-flex justify-content-between align-items-center mb-2">
                <span class="fw-bold text-info small">Resultado de la Operación:</span>
                <a id="btnDownloadWalk" href="#" class="btn btn-sm btn-success fw-bold" style="display: none;">📥 Descargar Walk Completo (.txt)</a>
            </div>
            <pre id="snmpResultBox" class="font-monospace small text-light mb-0" style="max-height: 350px; overflow-y: auto; white-space: pre-wrap; word-break: break-all;"></pre>
        </div>
    </div>
</div>

<div class="modal fade" id="deviceDetailModal" tabindex="-1" aria-hidden="true">
    <div class="modal-dialog modal-dialog-centered modal-xl">
        <div class="modal-content">
            <div class="modal-header bg-dark text-white">
                <h5 class="modal-title" id="modalDeviceTitle">📄 Detalles Completos de la ONT</h5>
                <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal" aria-label="Close"></button>
            </div>
            <div class="modal-body">
                <div class="row g-3">
                    <div class="col-md-4">
                        <div class="card-detail-section">
                            <h6>1. Hardware & Firmware</h6>
                            <div class="detail-item"><span class="detail-label">Fabricante:</span><span class="detail-val" id="m_manufacturer">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Modelo:</span><span class="detail-val" id="m_model">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Nro. Serie:</span><span class="detail-val" id="m_serial">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Versión HW:</span><span class="detail-val" id="m_hw_ver">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Versión FW:</span><span class="detail-val" id="m_firmware">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">RAM Total:</span><span class="detail-val" id="m_ram">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">RAM Libre:</span><span class="detail-val" id="m_ram_free">N/A</span></div>
                        </div>
                    </div>
                    <div class="col-md-4">
                        <div class="card-detail-section">
                            <h6>2. Direccionamiento IP</h6>
                            <div class="detail-item"><span class="detail-label">IP Datos:</span><span class="detail-val" id="m_ip_data">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">IP Gestión:</span><span class="detail-val" id="m_ip_mgmt">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">IP IPTV:</span><span class="detail-val" id="m_ip_iptv">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">IP VoIP:</span><span class="detail-val" id="m_ip_voip">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">DNS:</span><span class="detail-val" id="m_dns">N/A</span></div>
                        </div>
                    </div>
                    <div class="col-md-4">
                        <div class="card-detail-section">
                            <h6>3. Módulo Óptico GPON</h6>
                            <div class="detail-item"><span class="detail-label">Estado Óptico:</span><span class="detail-val" id="m_opt_status">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Potencia Tx:</span><span class="detail-val" id="m_opt_tx">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Potencia Rx:</span><span class="detail-val text-info" id="m_opt_rx">N/A</span></div>
                            <div class="detail-item"><span class="detail-label">Temperatura:</span><span class="detail-val" id="m_opt_temp">N/A</span></div>
                        </div>
                    </div>
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cerrar</button>
            </div>
        </div>
    </div>
</div>

<div class="modal fade" id="snmpDetailModal" tabindex="-1" aria-hidden="true">
    <div class="modal-dialog modal-dialog-centered modal-xl">
        <div class="modal-content">
            <div class="modal-header bg-dark text-white d-flex justify-content-between align-items-center">
                <h5 class="modal-title" id="modalSnmpTitle">📊 Servicios SNMP y OIDs</h5>
                <div>
                    <a id="btnExportCsv" href="#" class="btn btn-sm btn-success me-2">📥 Descargar CSV</a>
                    <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal" aria-label="Close"></button>
                </div>
            </div>
            <div class="modal-body p-0">
                <div class="table-responsive" style="max-height: 600px;">
                    <table class="table table-striped table-hover mb-0 align-middle">
                        <thead class="table-dark sticky-top">
                            <tr>
                                <th>Nombre del Servicio (Nagios)</th>
                                <th>OID SNMP</th>
                                <th>Valor Actual</th>
                            </tr>
                        </thead>
                        <tbody id="snmpServicesTableBody">
                            <tr><td colspan="3" class="text-center py-4">Cargando OIDs...</td></tr>
                        </tbody>
                    </table>
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cerrar</button>
            </div>
        </div>
    </div>
</div>

<div class="modal fade" id="metricsModal" tabindex="-1" aria-hidden="true">
    <div class="modal-dialog modal-dialog-centered modal-xl">
        <div class="modal-content">
            <div class="modal-header bg-dark text-white d-flex justify-content-between align-items-center">
                <h5 class="modal-title" id="metricsModalTitle">📈 Reporte de Estabilidad y Caídas</h5>
                <div>
                    <button id="btnToggleModalMonitoring" class="btn btn-sm btn-success me-2">▶️ Iniciar Monitoreo</button>
                    <a id="btnExportExcelMetrics" href="#" class="btn btn-sm btn-success fw-bold me-2">📥 Descargar Excel con Gráfico</a>
                    <button class="btn btn-sm btn-outline-warning me-2" onclick="reiniciarMonitoreoDispositivo()">🧹 Reiniciar Historial de este Equipo</button>
                    <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal" aria-label="Close"></button>
                </div>
            </div>
            <div class="modal-body">
                <div class="row g-3 mb-4">
                    <div class="col-md-4">
                        <div class="card p-3 text-center border-info">
                            <span class="text-secondary small">Índice de Estabilidad (SLA)</span>
                            <h2 id="sla_percentage" class="fw-bold text-info my-1">-- %</h2>
                            <span class="small text-muted">Disponibilidad en muestras</span>
                        </div>
                    </div>
                    <div class="col-md-4">
                        <div class="card p-3 text-center border-danger">
                            <span class="text-secondary small">Cantidad de Caídas</span>
                            <h2 id="total_drops_count" class="fw-bold text-danger my-1">0</h2>
                            <span class="small text-muted">Eventos registrados</span>
                        </div>
                    </div>
                    <div class="col-md-4">
                        <div class="card p-3 text-center border-warning">
                            <span class="text-secondary small">Total Muestras Analizadas</span>
                            <h2 id="total_samples_count" class="fw-bold text-warning my-1">0</h2>
                            <span class="small text-muted">Sondeos de red (cada 10s)</span>
                        </div>
                    </div>
                </div>

                <div class="card p-3 mb-4">
                    <h6 class="text-info fw-bold mb-3">Histórico Visual de Estabilidad (1 = Online / Verde, 0 = Offline / Rojo)</h6>
                    <div style="height: 250px; position: relative;">
                        <canvas id="stabilityChart"></canvas>
                    </div>
                </div>

                <h6 class="text-info fw-bold mb-2">📋 Historial de Caídas Detectadas</h6>
                <div class="table-responsive" style="max-height: 220px;">
                    <table class="table table-sm table-striped align-middle">
                        <thead class="table-dark">
                            <tr>
                                <th>#</th>
                                <th>Inicio Caída</th>
                                <th>Recuperación</th>
                                <th>Duración Fuera de Servicio</th>
                            </tr>
                        </thead>
                        <tbody id="dropsLogTableBody">
                            <tr><td colspan="4" class="text-center text-muted">Sin registros de caídas.</td></tr>
                        </tbody>
                    </table>
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cerrar</button>
            </div>
        </div>
    </div>
</div>

<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
<script>
    let currentDevices = [];
    let currentSnmpDevices = [];
    let myChartInstance = null;
    let currentModalHost = "";
    let activeMonitored = [];

    function cargarMonitoreados() {
        return fetch('/api/sentinel/monitoring/list')
            .then(res => res.json())
            .then(data => { activeMonitored = data || []; });
    }

    function iniciarMonitoreoTarget(target) {
        fetch('/api/sentinel/monitoring/start', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ target: target })
        })
        .then(res => res.json())
        .then(data => {
            cargarTodo();
        });
    }

    function detenerMonitoreoTarget(target) {
        fetch('/api/sentinel/monitoring/stop', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ target: target })
        })
        .then(res => res.json())
        .then(data => {
            cargarTodo();
        });
    }

    function cargarTR069() {
        fetch('/api/sentinel/tr069')
            .then(res => res.json())
            .then(data => {
                currentDevices = data;
                let total = data.length;
                let opticAlerts = 0;
                let offlineCount = 0;
                const tbody = document.getElementById('tr069TableBody');
                tbody.innerHTML = '';

                if (total === 0) {
                    tbody.innerHTML = '<tr><td colspan="9" class="text-center py-4">No hay ONTs registradas.</td></tr>';
                    return;
                }

                data.forEach(d => {
                    if (d.status !== 'Online') offlineCount++;
                    let rxVal = parseFloat(d.optical_rx);
                    let isRxBad = !isNaN(rxVal) && rxVal < -27.0;
                    if (isRxBad) opticAlerts++;

                    let isMon = activeMonitored.includes(d.serial_number);
                    let monBtn = isMon ? 
                        `<button class="btn btn-sm btn-danger me-1" onclick="detenerMonitoreoTarget('${d.serial_number}')">⏹️ Detener</button>` : 
                        `<button class="btn btn-sm btn-success me-1" onclick="iniciarMonitoreoTarget('${d.serial_number}')">▶️ Monitorear</button>`;

                    const tr = document.createElement('tr');
                    tr.innerHTML = `
                        <td><span class="badge ${d.status === 'Online' ? 'badge-success-custom' : 'badge-danger-custom'}">${d.status}</span></td>
                        <td class="fw-bold font-monospace">${d.serial_number}</td>
                        <td>${d.model_name || 'N/A'}</td>
                        <td>${d.hardware_version || 'N/A'}</td>
                        <td>${d.software_version || 'N/A'}</td>
                        <td>${d.ip_data || d.ip_address || 'N/A'}</td>
                        <td class="${isRxBad ? 'text-danger fw-bold' : ''}">${d.optical_rx || 'N/A'}</td>
                        <td class="small text-secondary">${d.last_inform || 'N/A'}</td>
                        <td class="text-end">
                            ${monBtn}
                            <button class="btn btn-sm btn-outline-info me-1" onclick="mostrarDetallesModal('${d.serial_number}')">Ficha</button>
                            <button class="btn btn-sm btn-info text-white" onclick="abrirModalEstabilidad('${d.serial_number}')">📊 Detalle</button>
                        </td>
                    `;
                    tbody.appendChild(tr);
                });

                document.getElementById('total_devs').textContent = total;
                document.getElementById('optic_alerts').textContent = opticAlerts;
                document.getElementById('offline_devs').textContent = offlineCount;
                actualizarTargetSelect();
            });
    }

    function cargarSNMP() {
        fetch('/api/sentinel/snmp')
            .then(res => res.json())
            .then(data => {
                currentSnmpDevices = data;
                const tbody = document.getElementById('snmpTableBody');
                tbody.innerHTML = '';

                if (data.length === 0) {
                    tbody.innerHTML = '<tr><td colspan="9" class="text-center py-4">No hay datos recolectados aún.</td></tr>';
                    return;
                }

                data.forEach(s => {
                    const isCm = s.ip_address.startsWith('172.16.');
                    let isMon = activeMonitored.includes(s.host_name);
                    let monBtn = isMon ? 
                        `<button class="btn btn-sm btn-danger me-1" onclick="detenerMonitoreoTarget('${s.host_name}')">⏹️ Detener</button>` : 
                        `<button class="btn btn-sm btn-success me-1" onclick="iniciarMonitoreoTarget('${s.host_name}')">▶️ Monitorear</button>`;

                    const tr = document.createElement('tr');
                    tr.innerHTML = `
                        <td><span class="badge ${s.status === 'Online' ? 'badge-success-custom' : 'badge-danger-custom'}">${s.status}</span></td>
                        <td class="fw-bold">${s.host_name}</td>
                        <td class="font-monospace">${s.ip_address}</td>
                        <td class="small text-truncate" style="max-width: 220px;" title="${s.sys_descr}">${s.sys_descr}</td>
                        <td>${s.serial_number}</td>
                        <td>${s.temperature}</td>
                        <td>${s.uptime}</td>
                        <td class="small text-secondary">${s.last_check}</td>
                        <td class="text-end">
                            ${monBtn}
                            ${isCm ? `<button class="btn btn-sm btn-outline-danger me-1" onclick="rebootCableModem('${s.ip_address}', '${s.host_name}')">⚡ Reboot</button>` : ''}
                            <button class="btn btn-sm btn-outline-warning me-1" onclick="abrirModalSnmpServices('${s.host_name}', '${s.ip_address}')">OIDs</button>
                            <button class="btn btn-sm btn-info text-white" onclick="abrirModalEstabilidad('${s.host_name}')">📊 Detalle</button>
                        </td>
                    `;
                    tbody.appendChild(tr);
                });
                actualizarTargetSelect();
            });
    }

    function actualizarTargetSelect() {
        const select = document.getElementById('snmpTargetSelect');
        const prevVal = select.value;
        select.innerHTML = '<option value="">-- Seleccionar Dispositivo --</option>';

        currentSnmpDevices.forEach(s => {
            select.innerHTML += `<option value="${s.ip_address}">${s.host_name} (${s.ip_address})</option>`;
        });

        currentDevices.forEach(d => {
            const ip = d.ip_data || d.ip_address;
            if (ip) {
                select.innerHTML += `<option value="${ip}">ONT ${d.serial_number} (${ip})</option>`;
            }
        });
        if (prevVal) select.value = prevVal;
    }

    function onOpTypeChange() {
        const op = document.getElementById('snmpOpType').value;
        const oidInput = document.getElementById('snmpOidInput');

        document.getElementById('snmpValContainer').style.display = (op === 'SET') ? 'block' : 'none';

        if (op === 'WALK') {
            if (!oidInput.value.trim() || oidInput.value === "1.3.6.1.2.1.1.3.0") oidInput.value = "1.3.6.1.2";
        }
    }

    function setOidPreset(oidVal) {
        document.getElementById('snmpOidInput').value = oidVal;
        document.getElementById('snmpOpType').value = 'WALK';
        onOpTypeChange();
    }

    function ejecutarSnmpCustom(e) {
        e.preventDefault();
        const ip = document.getElementById('snmpTargetSelect').value;
        const op = document.getElementById('snmpOpType').value;
        const oid = document.getElementById('snmpOidInput').value.trim();
        const val = document.getElementById('snmpValueInput').value.trim();

        const container = document.getElementById('snmpResultContainer');
        const box = document.getElementById('snmpResultBox');
        const btnDL = document.getElementById('btnDownloadWalk');

        if (!ip || !oid) return;

        container.style.display = 'block';
        btnDL.style.display = 'none';
        box.textContent = op === 'WALK' ? 'Ejecutando SNMP BULKWALK (descarga rápida optimizada)...' : 'Ejecutando comando SNMP...';

        fetch('/api/sentinel/snmp/exec', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ ip, op, oid, value: val })
        })
        .then(res => res.json())
        .then(data => {
            box.textContent = data.result || data.error;
            if (op === 'WALK' && !data.error) {
                btnDL.href = `/api/sentinel/snmp/walk/download?ip=${ip}&oid=${encodeURIComponent(oid)}`;
                btnDL.style.display = 'inline-block';
            }
        });
    }

    function rebootCableModem(ip, name) {
        if (!confirm(`¿Confirma el reinicio del Cable Módem ${name} (${ip})?`)) return;

        fetch('/api/sentinel/snmp/reboot', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ ip })
        })
        .then(res => res.json())
        .then(data => {
            alert(data.message || data.error);
        });
    }

    function reiniciarMonitoreoGlobal() {
        if (!confirm('¿Desea borrar el historial global de caídas y muestras registradas para TODOS los dispositivos?')) return;

        fetch('/api/sentinel/reset_monitoring', { method: 'POST' })
            .then(res => res.json())
            .then(data => {
                alert(data.message);
                cargarTodo();
            });
    }

    function reiniciarMonitoreoDispositivo() {
        if (!currentModalHost) return;
        if (!confirm(`¿Desea reiniciar el historial de muestras y caídas exclusivamente de "${currentModalHost}"?`)) return;

        fetch(`/api/sentinel/reset_monitoring/${encodeURIComponent(currentModalHost)}`, { method: 'POST' })
            .then(res => res.json())
            .then(data => {
                alert(data.message);
                abrirModalEstabilidad(currentModalHost);
            });
    }

    function abrirModalSnmpServices(hostName, ip) {
        document.getElementById('modalSnmpTitle').textContent = `📊 Servicios SNMP y OIDs — ${hostName} (${ip})`;
        document.getElementById('btnExportCsv').href = `/api/sentinel/snmp/services/export/${hostName}`;
        const tbody = document.getElementById('snmpServicesTableBody');
        tbody.innerHTML = '<tr><td colspan="3" class="text-center py-3">Cargando OIDs...</td></tr>';

        new bootstrap.Modal(document.getElementById('snmpDetailModal')).show();

        fetch(`/api/sentinel/snmp/services/${hostName}`)
            .then(res => res.json())
            .then(services => {
                tbody.innerHTML = '';
                if (services.length === 0) {
                    tbody.innerHTML = '<tr><td colspan="3" class="text-center py-3 text-secondary">Este equipo se monitorea por Ping (sin OIDs detallados).</td></tr>';
                    return;
                }
                services.forEach(svc => {
                    const tr = document.createElement('tr');
                    tr.innerHTML = `
                        <td class="fw-bold text-info">${svc.service_name}</td>
                        <td class="font-monospace text-secondary small">${svc.oid}</td>
                        <td class="font-monospace text-success fw-bold">${svc.value}</td>
                    `;
                    tbody.appendChild(tr);
                });
            });
    }

    function mostrarDetallesModal(serial) {
        const d = currentDevices.find(x => x.serial_number === serial);
        if (!d) return;
        document.getElementById('modalDeviceTitle').textContent = `📄 Detalles — ${d.model_name} (${d.serial_number})`;
        document.getElementById('m_manufacturer').textContent = d.manufacturer || 'Sagemcom';
        document.getElementById('m_model').textContent = d.model_name || 'N/A';
        document.getElementById('m_serial').textContent = d.serial_number || 'N/A';
        document.getElementById('m_hw_ver').textContent = d.hardware_version || 'N/A';
        document.getElementById('m_firmware').textContent = d.software_version || 'N/A';
        document.getElementById('m_ram').textContent = d.ram_total || 'N/A';
        document.getElementById('m_ram_free').textContent = d.ram_free || 'N/A';

        document.getElementById('m_ip_data').textContent = d.ip_data || d.ip_address || 'N/A';
        document.getElementById('m_ip_mgmt').textContent = d.ip_mgmt || 'N/A';
        document.getElementById('m_ip_iptv').textContent = d.ip_iptv || 'N/A';
        document.getElementById('m_ip_voip').textContent = d.ip_voip || 'N/A';
        document.getElementById('m_dns').textContent = d.dns_servers || 'N/A';

        document.getElementById('m_opt_status').textContent = d.optical_status || 'Up';
        document.getElementById('m_opt_tx').textContent = d.optical_tx || 'N/A';
        document.getElementById('m_opt_rx').textContent = d.optical_rx || 'N/A';
        document.getElementById('m_opt_temp').textContent = d.optical_temp || 'N/A';

        new bootstrap.Modal(document.getElementById('deviceDetailModal')).show();
    }

    function abrirModalEstabilidad(hostName) {
        currentModalHost = hostName;
        document.getElementById('metricsModalTitle').textContent = `📈 Reporte de Estabilidad — ${hostName}`;
        document.getElementById('btnExportExcelMetrics').href = `/api/sentinel/metrics/export_excel/${encodeURIComponent(hostName)}`;

        const btnMonModal = document.getElementById('btnToggleModalMonitoring');
        if (activeMonitored.includes(hostName)) {
            btnMonModal.className = "btn btn-sm btn-danger me-2";
            btnMonModal.innerHTML = "⏹️ Detener Monitoreo";
            btnMonModal.onclick = function() { detenerMonitoreoTarget(hostName); abrirModalEstabilidad(hostName); };
        } else {
            btnMonModal.className = "btn btn-sm btn-success me-2";
            btnMonModal.innerHTML = "▶️ Iniciar Monitoreo";
            btnMonModal.onclick = function() { iniciarMonitoreoTarget(hostName); abrirModalEstabilidad(hostName); };
        }

        new bootstrap.Modal(document.getElementById('metricsModal')).show();

        fetch(`/api/sentinel/metrics/${encodeURIComponent(hostName)}`)
            .then(res => res.json())
            .then(data => {
                document.getElementById('sla_percentage').textContent = `${data.sla_percentage}%`;
                document.getElementById('total_drops_count').textContent = data.total_drops;
                document.getElementById('total_samples_count').textContent = data.total_samples;

                const dropsTbody = document.getElementById('dropsLogTableBody');
                dropsTbody.innerHTML = '';

                if (data.drops.length === 0) {
                    dropsTbody.innerHTML = '<tr><td colspan="4" class="text-center text-success py-3"> Excelente: No se registraron caídas en el historial.</td></tr>';
                } else {
                    data.drops.forEach((drop, idx) => {
                        const tr = document.createElement('tr');
                        const duracion = drop.duration_seconds !== null ? `${drop.duration_seconds} seg` : '🔴 Caída Activa (En progreso...)';
                        tr.innerHTML = `
                            <td class="fw-bold">${idx + 1}</td>
                            <td class="text-danger fw-bold">${drop.drop_time}</td>
                            <td class="${drop.recovery_time ? 'text-success' : 'text-warning fw-bold'}">${drop.recovery_time || 'Aún Offline'}</td>
                            <td class="fw-bold">${duracion}</td>
                        `;
                        dropsTbody.appendChild(tr);
                    });
                }

                renderGraph(data.history);
            });
    }

    function renderGraph(historyData) {
        const ctx = document.getElementById('stabilityChart').getContext('2d');
        const labels = historyData.map(h => h.timestamp.split(' ')[1] || h.timestamp);
        const dataValues = historyData.map(h => h.status === 'Online' ? 1 : 0);

        if (myChartInstance) {
            myChartInstance.destroy();
        }

        myChartInstance = new Chart(ctx, {
            type: 'line',
            data: {
                labels: labels,
                datasets: [{
                    label: 'Estado',
                    data: dataValues,
                    borderWidth: 3,
                    stepped: true,
                    fill: true,
                    pointRadius: 4,
                    pointBackgroundColor: dataValues.map(v => v === 1 ? '#10b981' : '#ef4444'),
                    pointBorderColor: dataValues.map(v => v === 1 ? '#10b981' : '#ef4444'),
                    segment: {
                        borderColor: ctx => (ctx.p0.parsed.y === 0 || ctx.p1.parsed.y === 0) ? '#ef4444' : '#10b981',
                        backgroundColor: ctx => (ctx.p0.parsed.y === 0 || ctx.p1.parsed.y === 0) ? 'rgba(239, 68, 68, 0.2)' : 'rgba(16, 185, 129, 0.2)'
                    }
                }]
            },
            options: {
                responsive: true,
                maintainAspectRatio: false,
                scales: {
                    y: {
                        min: -0.1,
                        max: 1.1,
                        ticks: {
                            stepSize: 1,
                            callback: value => value === 1 ? 'Online' : (value === 0 ? 'Offline' : '')
                        },
                        grid: { color: '#334155' }
                    },
                    x: {
                        grid: { color: '#334155' }
                    }
                },
                plugins: {
                    legend: { display: false }
                }
            }
        });
    }

    async function cargarTodo() {
        await cargarMonitoreados();
        cargarTR069();
        cargarSNMP();
    }

    cargarTodo();
    setInterval(cargarTodo, 15000);
</script>
</body>
</html>
"""


@app.route('/')
def index():
    return render_template_string(HTML_DASHBOARD)


@app.route('/api/sentinel/monitoring/start', methods=['POST'])
def api_start_monitoring():
    data = request.get_json() or {}
    target = data.get("target")
    if not target:
        return jsonify({"error": "Target no especificado"}), 400
    with MONITORED_LOCK:
        MONITORED_TARGETS.add(target)

    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("INSERT OR IGNORE INTO monitored_targets (target_name) VALUES (?)", (target,))
    conn.commit()
    conn.close()

    return jsonify({"message": f"Monitoreo iniciado para '{target}'", "monitored": list(MONITORED_TARGETS)})


@app.route('/api/sentinel/monitoring/stop', methods=['POST'])
def api_stop_monitoring():
    data = request.get_json() or {}
    target = data.get("target")
    if not target:
        return jsonify({"error": "Target no especificado"}), 400
    with MONITORED_LOCK:
        MONITORED_TARGETS.discard(target)

    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM monitored_targets WHERE target_name = ?", (target,))
    conn.commit()
    conn.close()

    return jsonify({"message": f"Monitoreo detenido para '{target}'", "monitored": list(MONITORED_TARGETS)})


@app.route('/api/sentinel/monitoring/list', methods=['GET'])
def api_list_monitoring():
    with MONITORED_LOCK:
        return jsonify(list(MONITORED_TARGETS))


@app.route('/api/sentinel/tr069')
def api_tr069():
    conn = get_acs_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT serial_number, model_name, software_version, hardware_version, 
               status, last_inform, manufacturer, optical_rx, optical_tx, 
               optical_temp, optical_status, ram_total, ram_free, 
               ip_data, ip_mgmt, ip_iptv, ip_voip, dns_servers, ip_address
        FROM devices ORDER BY last_inform DESC
    """)
    rows = cursor.fetchall()
    conn.close()
    return jsonify([dict(row) for row in rows])


@app.route('/api/sentinel/snmp')
def api_snmp():
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT host_name, ip_address, sys_descr, uptime, serial_number, temperature, status, last_check
        FROM snmp_cache ORDER BY host_name ASC
    """)
    rows = cursor.fetchall()
    conn.close()
    return jsonify([dict(row) for row in rows])


@app.route('/api/sentinel/snmp/services/<host_name>')
def api_snmp_services(host_name):
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT service_name, oid, value, last_check
        FROM snmp_services_cache WHERE host_name = ? ORDER BY service_name ASC
    """, (host_name,))
    rows = cursor.fetchall()
    conn.close()
    return jsonify([dict(row) for row in rows])


@app.route('/api/sentinel/snmp/services/export/<host_name>')
def export_snmp_services_csv(host_name):
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT service_name, oid, value, last_check
        FROM snmp_services_cache WHERE host_name = ? ORDER BY service_name ASC
    """, (host_name,))
    rows = cursor.fetchall()
    conn.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Servicio", "OID SNMP", "Valor", "Ultima Consulta"])
    for r in rows:
        writer.writerow([r["service_name"], r["oid"], r["value"], r["last_check"]])

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-disposition": f"attachment; filename=OIDs_{host_name}.csv"}
    )


@app.route('/api/sentinel/snmp/reboot', methods=['POST'])
def api_snmp_reboot():
    data = request.get_json() or {}
    ip = data.get("ip")
    if not ip:
        return jsonify({"error": "IP no especificada"}), 400

    cmd = ["snmpset", "-v2c", "-c", "private", ip, "1.3.6.1.2.1.69.1.1.3.0", "i", "1"]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=5)
        if res.returncode == 0:
            return jsonify({"message": f"Orden de reinicio enviada exitosamente a {ip}"})
        else:
            return jsonify({"error": f"Fallo al enviar SNMP SET: {res.stderr.strip() or res.stdout.strip()}"}), 500
    except Exception as e:
        return jsonify({"error": f"Excepción ejecutando snmpset: {str(e)}"}), 500


@app.route('/api/sentinel/snmp/exec', methods=['POST'])
def api_sentinel_exec():
    global GLOBAL_WALK_CACHE
    data = request.get_json() or {}
    ip = data.get("ip")
    op = data.get("op", "GET")
    oid = data.get("oid", "1.3.6.1.2").strip() or "1.3.6.1.2"
    val_str = data.get("value", "").strip()

    if not ip:
        return jsonify({"error": "Dirección IP faltante"}), 400

    try:
        if op == "WALK":
            cmd = ["snmpbulkwalk", "-v2c", "-c", "private", "-t", "5", "-r", "3", ip, oid]
            try:
                res = subprocess.run(cmd, capture_output=True, timeout=60)

                raw_out = res.stdout if res.stdout else res.stderr
                if isinstance(raw_out, bytes):
                    raw_out = raw_out.decode('utf-8', errors='ignore')

                # Guardar en caché el resultado completo para la descarga instantánea
                GLOBAL_WALK_CACHE = {"ip": ip, "oid": oid, "content": raw_out}

                out_lines = raw_out.strip().splitlines()
                total_lines = len(out_lines)
                preview = "\n".join(out_lines[:300])
                if total_lines > 300:
                    preview += f"\n\n... [ Se encontraron {total_lines} OIDs en total. Haz clic en el botón de descarga arriba para obtener el archivo completo ]"
                return jsonify({"result": preview if res.returncode == 0 else f"Error: {preview}"})

            except subprocess.TimeoutExpired as e:
                partial_output = e.stdout if e.stdout else e.stderr
                if isinstance(partial_output, bytes):
                    partial_output = partial_output.decode('utf-8', errors='ignore')

                # Guardar también el parcial en caché si hubo timeout
                GLOBAL_WALK_CACHE = {"ip": ip, "oid": oid, "content": partial_output}

                if partial_output:
                    out_lines = partial_output.strip().splitlines()
                    preview = "\n".join(out_lines[:300])
                    preview = f"⚠️ TIMEOUT Alcanzado (60s): Se muestra el resultado parcial acumulado ({len(out_lines)} líneas):\n\n" + preview
                    return jsonify({"result": preview})
                else:
                    return jsonify({
                        "result": "⚠️ El comando SNMP BULKWALK superó el tiempo límite de 60s sin devolver datos iniciales."})

        elif op == "SET":
            if not val_str:
                return jsonify({"error": "Valor faltante para operación SET (ej: i 1)"}), 400
            val_parts = val_str.split(" ", 1)
            if len(val_parts) < 2:
                return jsonify({"error": "El valor SET debe incluir tipo y valor (ej: i 1 o s mi_texto)"}), 400
            cmd = ["snmpset", "-v2c", "-c", "private", ip, oid, val_parts[0], val_parts[1]]
            res = subprocess.run(cmd, capture_output=True, timeout=5)
            out = res.stdout if res.stdout else res.stderr
            if isinstance(out, bytes):
                out = out.decode('utf-8', errors='ignore')
            out = out.strip()
            return jsonify({"result": out if res.returncode == 0 else f"Error: {out}"})

        else:  # GET
            cmd = ["snmpget", "-v2c", "-c", "private", "-t", "3", ip, oid]
            res = subprocess.run(cmd, capture_output=True, timeout=5)
            out = res.stdout if res.stdout else res.stderr
            if isinstance(out, bytes):
                out = out.decode('utf-8', errors='ignore')
            out = out.strip()
            return jsonify({"result": out if res.returncode == 0 else f"Error: {out}"})

    except Exception as e:
        return jsonify({"error": f"Error ejecutando comando: {str(e)}"}), 500


@app.route('/api/sentinel/snmp/walk/download')
def download_snmp_walk():
    global GLOBAL_WALK_CACHE
    ip = request.args.get("ip")
    oid = request.args.get("oid", "1.3.6.1.2").strip() or "1.3.6.1.2"
    if not ip:
        return "IP faltante", 400

    # Si coincide con la última ejecución en caché, se entrega de inmediato SIN re-ejecutar nada
    if GLOBAL_WALK_CACHE["ip"] == ip and GLOBAL_WALK_CACHE["oid"] == oid and GLOBAL_WALK_CACHE["content"]:
        content = GLOBAL_WALK_CACHE["content"]
    else:
        # Fallback por si acaso se descarga directamente sin pasar por la consola
        cmd = ["snmpbulkwalk", "-v2c", "-c", "private", "-t", "5", "-r", "3", ip, oid]
        try:
            res = subprocess.run(cmd, capture_output=True, timeout=90)
            raw_content = res.stdout if res.stdout else res.stderr
            if isinstance(raw_content, bytes):
                content = raw_content.decode('utf-8', errors='ignore')
            else:
                content = raw_content
        except Exception as e:
            content = f"Excepción ejecutando snmpbulkwalk: {str(e)}"

    clean_oid = oid.replace(".", "_")
    filename = f"snmpwalk_{ip}_oid_{clean_oid}.txt"

    return Response(
        content,
        mimetype="text/plain",
        headers={"Content-disposition": f"attachment; filename={filename}"}
    )


@app.route('/api/sentinel/reset_monitoring', methods=['POST'])
def api_reset_monitoring_global():
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM device_history")
    cursor.execute("DELETE FROM device_drops")
    conn.commit()
    conn.close()
    return jsonify({"message": "Histórico global de estabilidad y caídas reiniciado correctamente."})


@app.route('/api/sentinel/reset_monitoring/<host_name>', methods=['POST'])
def api_reset_monitoring_device(host_name):
    conn = get_snmp_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM device_history WHERE host_name = ?", (host_name,))
    cursor.execute("DELETE FROM device_drops WHERE host_name = ?", (host_name,))
    conn.commit()
    conn.close()
    return jsonify({"message": f"Histórico de estabilidad y caídas de '{host_name}' reiniciado correctamente."})


@app.route('/api/sentinel/metrics/<host_name>')
def api_sentinel_metrics(host_name):
    conn = get_snmp_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT status, timestamp FROM device_history
        WHERE host_name = ? ORDER BY id DESC LIMIT 50000
    """, (host_name,))
    history = [dict(row) for row in cursor.fetchall()]
    history.reverse()

    cursor.execute("SELECT COUNT(*) as total FROM device_history WHERE host_name = ?", (host_name,))
    total_samples = cursor.fetchone()["total"]

    cursor.execute("SELECT COUNT(*) as online_cnt FROM device_history WHERE host_name = ? AND status = 'Online'",
                   (host_name,))
    online_samples = cursor.fetchone()["online_cnt"]

    if total_samples > 0:
        sla_percentage = round((online_samples / total_samples) * 100, 2)
    else:
        sla_percentage = 0.0

    cursor.execute("""
        SELECT drop_time, recovery_time, duration_seconds
        FROM device_drops WHERE host_name = ? ORDER BY id DESC LIMIT 2000
    """, (host_name,))
    drops = [dict(row) for row in cursor.fetchall()]

    conn.close()

    return jsonify({
        "sla_percentage": sla_percentage,
        "total_samples": total_samples,
        "total_drops": len(drops),
        "history": history,
        "drops": drops
    })


@app.route('/api/sentinel/metrics/export_excel/<host_name>')
def export_metrics_excel(host_name):
    conn = get_snmp_db()
    cursor = conn.cursor()

    # Obtener TODO el historial sin límite
    cursor.execute("""
        SELECT status, timestamp FROM device_history
        WHERE host_name = ? ORDER BY id ASC
    """, (host_name,))
    full_history = [dict(row) for row in cursor.fetchall()]

    cursor.execute("SELECT COUNT(*) as total FROM device_history WHERE host_name = ?", (host_name,))
    total_samples = cursor.fetchone()["total"]

    cursor.execute("SELECT COUNT(*) as online_cnt FROM device_history WHERE host_name = ? AND status = 'Online'",
                   (host_name,))
    online_samples = cursor.fetchone()["online_cnt"]

    sla_percentage = round((online_samples / total_samples) * 100, 2) if total_samples > 0 else 0.0

    # Obtener TODAS las caídas sin límite
    cursor.execute("""
        SELECT drop_time, recovery_time, duration_seconds
        FROM device_drops WHERE host_name = ? ORDER BY id DESC
    """, (host_name,))
    full_drops = [dict(row) for row in cursor.fetchall()]
    conn.close()

    # Generar gráfico con Matplotlib idéntico al Dashboard
    fig, ax = plt.subplots(figsize=(10, 3.8), facecolor='#0f172a')
    ax.set_facecolor('#1e293b')

    timestamps = [h['timestamp'].split(' ')[1] if ' ' in h['timestamp'] else h['timestamp'] for h in full_history]
    y_vals = [1 if h['status'] == 'Online' else 0 for h in full_history]

    if timestamps:
        ax.step(range(len(timestamps)), y_vals, where='post', color='#10b981', linewidth=2)
        colors = ['#10b981' if v == 1 else '#ef4444' for v in y_vals]
        ax.scatter(range(len(timestamps)), y_vals, color=colors, zorder=5, s=20)

    ax.set_yticks([0, 1])
    ax.set_yticklabels(['Offline', 'Online'], color='#f8fafc', fontsize=9, fontweight='bold')
    ax.tick_params(axis='x', colors='#f8fafc', labelsize=8)
    ax.tick_params(axis='y', colors='#f8fafc')
    ax.set_title(f"Histórico Visual de Estabilidad — {host_name}", color='#38bdf8', fontsize=11, fontweight='bold')
    ax.grid(True, color='#334155', linestyle='--', linewidth=0.5)

    if len(timestamps) > 20:
        step = max(1, len(timestamps) // 15)
        ax.set_xticks(range(0, len(timestamps), step))
        ax.set_xticklabels([timestamps[i] for i in range(0, len(timestamps), step)], rotation=45, ha='right')
    elif len(timestamps) > 0:
        ax.set_xticks(range(len(timestamps)))
        ax.set_xticklabels(timestamps, rotation=45, ha='right')

    plt.tight_layout()
    img_buf = io.BytesIO()
    plt.savefig(img_buf, format='png', dpi=120)
    plt.close(fig)
    img_buf.seek(0)

    # Crear Libro de Excel con OpenPyXL
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Estabilidad"
    ws.views.sheetView[0].showGridLines = True

    # Estilos de formato
    title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    metric_fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')

    ws.merge_cells('A1:D1')
    ws['A1'] = f"Reporte de Estabilidad y Caídas — {host_name}"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='38BDF8')
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    # Tarjetas Métricas
    ws['A3'] = "Índice de Estabilidad (SLA)"
    ws['A4'] = f"{sla_percentage}%"
    ws['B3'] = "Cantidad de Caídas"
    ws['B4'] = len(full_drops)
    ws['C3'] = "Total Muestras Analizadas"
    ws['C4'] = total_samples

    for col in ['A', 'B', 'C']:
        ws[f'{col}3'].font = Font(bold=True, color='94A3B8', size=9)
        ws[f'{col}3'].fill = metric_fill
        ws[f'{col}3'].alignment = Alignment(horizontal='center')
        ws[f'{col}4'].font = Font(bold=True, size=13, color='FFFFFF')
        ws[f'{col}4'].fill = title_fill
        ws[f'{col}4'].alignment = Alignment(horizontal='center')

    # Incrustar Gráfico al costado
    img = OpenpyxlImage(img_buf)
    img.width = 620
    img.height = 240
    ws.add_image(img, 'F1')

    # Tabla 1: Historial de Caídas
    ws['A6'] = "#"
    ws['B6'] = "Inicio Caída"
    ws['C6'] = "Recuperación"
    ws['D6'] = "Duración Fuera de Servicio"

    for col in ['A6', 'B6', 'C6', 'D6']:
        ws[col].font = Font(bold=True, color='38BDF8')
        ws[col].fill = header_fill

    row_idx = 7
    if not full_drops:
        ws[f'A{row_idx}'] = 1
        ws[f'B{row_idx}'] = "Sin caídas registradas"
        ws[f'C{row_idx}'] = "N/A"
        ws[f'D{row_idx}'] = "0 seg"
        row_idx += 1
    else:
        for idx, drop in enumerate(full_drops, 1):
            dur = f"{drop['duration_seconds']} seg" if drop[
                                                           'duration_seconds'] is not None else "Caída Activa (En progreso...)"
            rec = drop['recovery_time'] if drop['recovery_time'] else "Aún Offline"
            ws[f'A{row_idx}'] = idx
            ws[f'B{row_idx}'] = drop['drop_time']
            ws[f'C{row_idx}'] = rec
            ws[f'D{row_idx}'] = dur
            row_idx += 1

    # Tabla 2: Histórico Completo de Muestras
    row_idx += 2
    ws.cell(row=row_idx, column=1, value="Histórico Completo de Muestras Analizadas").font = Font(bold=True, size=11,
                                                                                                  color='38BDF8')
    row_idx += 1

    ws.cell(row=row_idx, column=1, value="Timestamp").font = Font(bold=True, color='FFFFFF')
    ws.cell(row=row_idx, column=1).fill = header_fill
    ws.cell(row=row_idx, column=2, value="Estado").font = Font(bold=True, color='FFFFFF')
    ws.cell(row=row_idx, column=2).fill = header_fill

    row_idx += 1
    for h in full_history:
        ws.cell(row=row_idx, column=1, value=h['timestamp'])
        st_cell = ws.cell(row=row_idx, column=2, value=h['status'])
        if h['status'] == 'Online':
            st_cell.font = Font(color='10B981', bold=True)
        else:
            st_cell.font = Font(color='EF4444', bold=True)
        row_idx += 1

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    excel_buf.seek(0)

    return Response(
        excel_buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-disposition": f"attachment; filename=Reporte_Estabilidad_{host_name}.xlsx"}
    )


if __name__ == '__main__':
    print("Iniciando CPE Sentinel Web (con Caché Instantánea de Walk) en puerto 9035...")
    app.run(host='0.0.0.0', port=9035, debug=False, threaded=True)